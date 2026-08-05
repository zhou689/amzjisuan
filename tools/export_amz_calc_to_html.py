from __future__ import annotations

import json
import re
from copy import copy
from html import escape
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


SOURCE = Path(r"E:/球类/AMZ运营计算器2.0.xlsx")
FALLBACK_SOURCE = Path("AMZ_calc_source.xlsx")
OUTPUT = Path("AMZ运营计算器2.0.html")
INDEX_OUTPUT = Path("index.html")
CELL_REF_RE = re.compile(r"(?<![A-Z0-9_])\$?([A-Z]{1,3})\$?(\d+)\b")
EMPHASIS_RESULT_LABELS = {"最低售价", "实际利润", "销售回款", "实际成本"}


def color_to_hex(color: Any) -> str | None:
    if not color or color.type is None:
        return None
    if color.type == "rgb" and color.rgb:
        rgb = color.rgb[-6:]
        if rgb == "000000" and getattr(color, "rgb", "").startswith("00"):
            return None
        return f"#{rgb}"
    if color.type == "indexed" and color.indexed is not None:
        try:
            value = COLOR_INDEX[color.indexed]
            if value:
                return f"#{value[-6:]}"
        except Exception:
            return None
    return None


def fill_to_hex(fill: PatternFill) -> str | None:
    if not fill or fill.fill_type is None:
        return None
    return color_to_hex(fill.fgColor) or color_to_hex(fill.start_color)


def border_style(side: Any) -> str:
    if not side or not side.style:
        return "1px solid transparent"
    color = color_to_hex(side.color) or "#a9b3c1"
    widths = {
        "hair": "1px",
        "thin": "1px",
        "medium": "2px",
        "thick": "3px",
        "double": "3px",
        "dashed": "1px",
        "dotted": "1px",
    }
    css_style = "solid"
    if side.style == "dashed":
        css_style = "dashed"
    elif side.style == "dotted":
        css_style = "dotted"
    elif side.style == "double":
        css_style = "double"
    return f"{widths.get(side.style, '1px')} {css_style} {color}"


def alignment_css(cell: Any) -> list[str]:
    styles: list[str] = []
    horizontal = cell.alignment.horizontal
    vertical = cell.alignment.vertical
    if horizontal:
        align = {"centerContinuous": "center", "general": "left"}.get(horizontal, horizontal)
        styles.append(f"text-align:{align}")
    if vertical:
        styles.append(f"vertical-align:{vertical}")
    if cell.alignment.wrap_text:
        styles.append("white-space:pre-wrap")
    return styles


def font_css(cell: Any) -> list[str]:
    styles: list[str] = []
    font = cell.font
    if font.name:
        styles.append(f"font-family:'{font.name}', 'Microsoft YaHei', Arial, sans-serif")
    if font.bold:
        styles.append("font-weight:700")
    if font.italic:
        styles.append("font-style:italic")
    if font.underline:
        styles.append("text-decoration:underline")
    return styles


def cell_css(cell: Any) -> str:
    styles = []
    styles.extend(alignment_css(cell))
    styles.extend(font_css(cell))
    return ";".join(styles)


def is_editable(cell: Any, referenced_blanks: set[str]) -> bool:
    if cell.value is None:
        return cell.coordinate in referenced_blanks
    if isinstance(cell.value, str):
        return False
    # Numeric constants are the natural spreadsheet inputs. Formula outputs stay locked.
    return not (isinstance(cell.value, str) and cell.value.startswith("="))


def raw_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, str, bool)):
        return value
    return str(value)


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if abs(value) >= 1000:
            text = f"{value:,.2f}"
        else:
            text = f"{value:.4f}"
        text = text.rstrip("0").rstrip(".")
        return text
    return str(value)


def column_roles(ws: Any, select_cells: dict[str, dict[str, str]], referenced_blanks: set[str]) -> list[str]:
    roles = ["spacer"] * ws.max_column
    unit_values = {"￥", "¥", "$", "%", "次", "单", "Kg", "kg", "月", "个", "¢", ".98"}
    for col in range(1, ws.max_column + 1):
        has_value = False
        has_unit = False
        has_label = False
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row, col)
            value = cell.value
            if cell.coordinate in select_cells or is_editable(cell, referenced_blanks):
                has_value = True
            elif isinstance(value, str) and value.startswith("="):
                has_value = True
            elif isinstance(value, (int, float)):
                has_value = True
            elif isinstance(value, str) and value.strip() in unit_values:
                has_unit = True
            elif value is not None:
                has_label = True
        if has_value:
            roles[col - 1] = "value"
        elif has_unit:
            roles[col - 1] = "unit"
        elif has_label:
            roles[col - 1] = "label"
    return roles


def build_workbook_data() -> dict[str, Any]:
    source = SOURCE if SOURCE.exists() else FALLBACK_SOURCE
    wb_formula = load_workbook(source, data_only=False)
    wb_values = load_workbook(source, data_only=True)
    sheets: list[dict[str, Any]] = []

    for ws in wb_formula.worksheets:
        ws_values = wb_values[ws.title]
        referenced_blanks: set[str] = set()
        for row in ws.iter_rows():
            for formula_cell in row:
                if isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                    for col, row_num in CELL_REF_RE.findall(formula_cell.value):
                        coord = f"{col}{row_num}"
                        if ws[coord].value is None:
                            referenced_blanks.add(coord)

        covered: set[tuple[int, int]] = set()
        merge_lookup: dict[tuple[int, int], tuple[int, int]] = {}
        for merged in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(merged))
            merge_lookup[(min_row, min_col)] = (max_row - min_row + 1, max_col - min_col + 1)
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    if (row, col) != (min_row, min_col):
                        covered.add((row, col))

        cells: list[dict[str, Any]] = []
        formulas: dict[str, str] = {}
        values: dict[str, Any] = {}
        select_cells: dict[str, dict[str, str]] = {}
        if ws.title in {"FBA最低售价及利润", "欧洲-FBA最低售价及利润"}:
            select_cells["H3"] = {
                "source": "E3",
                "zeroLabel": "0",
                "syncLabel": "同步广告点击",
            }

        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                if (row, col) in covered:
                    continue
                cell = ws.cell(row, col)
                cached = ws_values[cell.coordinate].value
                value = raw_value(cell.value)
                formula = value if isinstance(value, str) and value.startswith("=") else None
                if formula:
                    formulas[cell.coordinate] = formula
                    values[cell.coordinate] = raw_value(cached)
                elif value is not None:
                    values[cell.coordinate] = value

                rowspan, colspan = merge_lookup.get((row, col), (1, 1))
                cells.append(
                    {
                        "coord": cell.coordinate,
                        "row": row,
                        "col": col,
                        "rowspan": rowspan,
                        "colspan": colspan,
                        "value": value,
                        "cached": raw_value(cached),
                        "display": display_value(cached if formula else value),
                        "formula": formula,
                        "editable": is_editable(cell, referenced_blanks),
                        "emphasis": bool(
                            formula
                            and col > 1
                            and str(ws.cell(row, col - 1).value or "").strip() in EMPHASIS_RESULT_LABELS
                        ),
                        "style": cell_css(cell),
                        "numberFormat": cell.number_format,
                    }
                )

        roles = column_roles(ws, select_cells, referenced_blanks)
        role_widths = {"label": 118, "value": 120, "unit": 30, "spacer": 36}
        col_widths = []
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            hidden = bool(ws.column_dimensions[letter].hidden)
            role = roles[col - 1]
            col_widths.append({"letter": letter, "px": role_widths[role], "hidden": hidden, "role": role})

        row_heights = []
        for row in range(1, ws.max_row + 1):
            height = ws.row_dimensions[row].height or 20
            hidden = bool(ws.row_dimensions[row].hidden)
            row_heights.append({"row": row, "px": round(height * 1.34), "hidden": hidden})

        sheets.append(
            {
                "name": ws.title,
                "rows": ws.max_row,
                "cols": ws.max_column,
                "cells": cells,
                "values": values,
                "formulas": formulas,
                "selectCells": select_cells,
                "colWidths": col_widths,
                "rowHeights": row_heights,
            }
        )

    return {"sheets": sheets}


def html_template(data: dict[str, Any]) -> str:
    workbook_json = json.dumps(data, ensure_ascii=False)
    display_names = {"尺寸": "单位互换"}
    sheet_tabs = "\n".join(
        f'<button class="tab{" is-active" if i == 0 else ""}" type="button" data-sheet="{i}">{escape(display_names.get(sheet["name"], sheet["name"]))}</button>'
        for i, sheet in enumerate(data["sheets"])
    )
    generated_at = "2026-08-05"
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>AMZ运营计算器2.0</title>
  <style>
    :root {{
      --bg: #eef3f8;
      --sidebar: #172235;
      --sidebar-2: #21304a;
      --panel: #ffffff;
      --panel-soft: #f8fafc;
      --text: #17202e;
      --muted: #66758b;
      --line: #d8e1ec;
      --line-strong: #b9c6d6;
      --accent: #1768e5;
      --accent-dark: #0f55be;
      --accent-soft: #e8f1ff;
      --input: #ffffff;
      --input-line: #d1a30f;
      --formula: #eef4fb;
      --danger: #c3352b;
      --shadow: 0 18px 50px rgba(23, 34, 53, .14);
      --radius: 8px;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font-family: "Microsoft YaHei", "PingFang SC", Arial, sans-serif;
      line-height: 1.4;
    }}
    .app-shell {{
      min-height: 100vh;
      display: grid;
      grid-template-columns: 248px minmax(0, 1fr);
    }}
    .sidebar {{
      position: sticky;
      top: 0;
      height: 100vh;
      display: flex;
      flex-direction: column;
      padding: 22px 16px;
      background: linear-gradient(180deg, var(--sidebar), var(--sidebar-2));
      color: #eef5ff;
    }}
    .brand {{
      padding: 0 8px 18px;
      border-bottom: 1px solid rgba(255, 255, 255, .12);
    }}
    .brand h1 {{
      margin: 0;
      font-size: 20px;
      letter-spacing: 0;
      line-height: 1.2;
    }}
    .brand p {{
      margin: 8px 0 0;
      color: rgba(238, 245, 255, .72);
      font-size: 14px;
    }}
    .tabs {{
      display: flex;
      flex-direction: column;
      gap: 7px;
      overflow-y: auto;
      padding: 18px 0;
    }}
    .tab {{
      width: 100%;
      min-height: 38px;
      justify-content: flex-start;
      border: 1px solid transparent;
      background: transparent;
      color: rgba(238, 245, 255, .76);
      text-align: left;
      border-radius: 7px;
      padding: 10px 10px;
      font-size: 14px;
    }}
    .tab:hover {{
      border-color: rgba(255, 255, 255, .14);
      background: rgba(255, 255, 255, .07);
      color: #fff;
    }}
    .tab.is-active {{
      color: #17202e;
      background: #ffffff;
      border-color: #ffffff;
      box-shadow: 0 8px 24px rgba(0, 0, 0, .16);
    }}
    .sidebar-note {{
      margin-top: auto;
      padding: 12px 10px 0;
      color: rgba(238, 245, 255, .62);
      font-size: 13px;
      border-top: 1px solid rgba(255, 255, 255, .12);
    }}
    .main-panel {{
      min-width: 0;
      padding: 22px 24px 30px;
    }}
    .topbar {{
      display: flex;
      align-items: flex-start;
      justify-content: space-between;
      gap: 18px;
      margin-bottom: 16px;
    }}
    .title-block h2 {{
      margin: 0;
      font-size: 22px;
      line-height: 1.22;
      letter-spacing: 0;
    }}
    .title-block p {{
      margin: 6px 0 0;
      color: var(--muted);
      font-size: 14px;
    }}
    .actions {{
      display: flex;
      align-items: center;
      gap: 8px;
      flex-wrap: wrap;
      justify-content: flex-end;
    }}
    button {{
      border: 1px solid var(--line);
      background: var(--panel);
      color: var(--text);
      border-radius: 6px;
      min-height: 40px;
      padding: 8px 12px;
      font: 600 14px/1 "Microsoft YaHei", Arial, sans-serif;
      cursor: pointer;
      transition: border-color .15s ease, background .15s ease, transform .15s ease, box-shadow .15s ease;
    }}
    button:hover {{
      border-color: var(--line-strong);
      transform: translateY(-1px);
    }}
    .primary {{
      background: var(--accent);
      border-color: var(--accent);
      color: #fff;
      box-shadow: 0 8px 20px rgba(23, 104, 229, .22);
    }}
    .primary:hover {{ background: var(--accent-dark); }}
    .summary {{
      display: grid;
      grid-template-columns: repeat(4, minmax(130px, 1fr));
      gap: 10px;
      margin-bottom: 14px;
    }}
    .metric {{
      min-height: 64px;
      padding: 11px 12px;
      border: 1px solid var(--line);
      border-radius: var(--radius);
      background: rgba(255, 255, 255, .78);
      box-shadow: 0 8px 22px rgba(23, 34, 53, .06);
    }}
    .metric span {{
      display: block;
      color: var(--muted);
      font-size: 13px;
    }}
    .metric strong {{
      display: block;
      margin-top: 4px;
      font-size: 20px;
      line-height: 1.1;
    }}
    .metric strong.small {{
      font-size: 15px;
      white-space: nowrap;
      overflow: hidden;
      text-overflow: ellipsis;
    }}
    .sheet-shell {{
      overflow: auto;
      border: 1px solid var(--line);
      background: var(--panel);
      box-shadow: var(--shadow);
      max-height: calc(100vh - 164px);
      border-radius: var(--radius);
    }}
    table.sheet {{
      border-collapse: separate;
      border-spacing: 0 10px;
      table-layout: fixed;
      width: max-content;
      min-width: 100%;
      background: #fff;
      padding: 18px 22px 28px;
    }}
    .sheet td {{
      min-width: 0;
      height: 42px;
      padding: 7px 6px;
      border: 0 !important;
      background: transparent !important;
      overflow: hidden;
      text-overflow: ellipsis;
      font-size: 15px;
      line-height: 1.34;
      white-space: nowrap;
    }}
    .sheet td.wrap {{ white-space: pre-wrap; }}
    .sheet td.role-label {{
      padding-right: 8px;
    }}
    .sheet td.role-value {{
      padding-left: 4px;
      padding-right: 4px;
    }}
    .sheet td.role-unit {{
      padding-left: 2px;
      padding-right: 4px;
      text-align: left !important;
    }}
    .sheet td.role-spacer {{
      padding-left: 0;
      padding-right: 0;
    }}
    .sheet td.blank {{
      padding-left: 0;
      padding-right: 0;
    }}
    .calculator-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(286px, 340px));
      align-items: start;
      justify-content: start;
      gap: 20px 22px;
      min-width: 920px;
      padding: 28px;
      background: #fff;
    }}
    .calculator-grid.fba-layout {{
      grid-template-columns: 340px 340px 340px;
      grid-template-areas:
        "cost ads traffic"
        "sales profit settings"
        "note note note";
      gap: 22px 28px;
      min-width: 1080px;
    }}
    .calc-card {{
      align-self: start;
      padding: 18px 18px 16px;
      border: 1px solid #e0e7f0;
      border-radius: var(--radius);
      background: #fbfdff;
    }}
    .calc-card h3 {{
      margin: 0 0 14px;
      font-size: 16px;
      line-height: 1.2;
      letter-spacing: 0;
    }}
    .calc-card.wide {{
      grid-column: span 2;
    }}
    .area-cost {{ grid-area: cost; }}
    .area-ads {{ grid-area: ads; }}
    .area-traffic {{ grid-area: traffic; }}
    .area-sales {{ grid-area: sales; }}
    .area-profit {{ grid-area: profit; }}
    .area-settings {{ grid-area: settings; }}
    .field-row {{
      display: grid;
      grid-template-columns: 110px 108px 28px;
      align-items: center;
      column-gap: 8px;
      min-height: 48px;
    }}
    .field-row.compact {{
      grid-template-columns: 110px 108px;
    }}
    .card-divider {{
      margin: 14px 0 12px;
      border: 0;
      border-top: 1px solid #e1e9f2;
    }}
    .subhead {{
      margin: 0 0 8px;
      font-size: 15px;
      font-weight: 700;
      color: #34445b;
    }}
    .field-label {{
      justify-self: end;
      font-size: 15px;
      font-weight: 700;
      color: var(--text);
      white-space: nowrap;
    }}
    .field-unit {{
      justify-self: start;
      font-size: 15px;
      color: var(--text);
      white-space: nowrap;
    }}
    .calc-note {{
      grid-column: 1 / -1;
      margin: 4px 0 0;
      padding-top: 14px;
      border-top: 1px solid #e6edf5;
      color: var(--muted);
      font-size: 14px;
      text-align: center;
      white-space: pre-wrap;
    }}
    .cell-input {{
      display: block;
      width: 100%;
      min-width: 108px;
      max-width: 108px;
      height: 40px;
      border: 0;
      outline: none;
      min-height: 40px;
      padding: 6px 9px;
      margin: 0 auto;
      background: var(--input);
      color: inherit;
      font: inherit;
      text-align: center;
      border-radius: 5px;
      box-shadow: inset 0 0 0 1px #d9e2ed, inset 0 -2px 0 var(--input-line);
    }}
    .cell-input:focus {{
      background: #fff;
      box-shadow: inset 0 0 0 1px var(--accent), inset 0 -2px 0 var(--accent);
    }}
    td.editable {{
      background: transparent !important;
      box-shadow: none;
    }}
    td.formula {{
      background: transparent !important;
    }}
    td.readonly-value {{
      background: #f6f8fb !important;
      border-radius: 5px;
      box-shadow: inset 0 0 0 1px #e0e7f0;
    }}
    .cell-value {{
      display: flex;
      align-items: center;
      justify-content: center;
      width: 108px;
      min-width: 108px;
      height: 40px;
      min-height: 40px;
      padding: 6px 9px;
      margin: 0 auto;
      border-radius: 5px;
      background: var(--formula);
      box-shadow: inset 0 0 0 1px #d9e4ef;
      font: inherit;
      font-variant-numeric: tabular-nums;
      text-align: center;
    }}
    td.numeric {{
      text-align: center !important;
      font-variant-numeric: tabular-nums;
    }}
    td.negative,
    .cell-input.negative,
    .cell-select.negative,
    .cell-value.negative,
    .cell-value.emphasis {{
      color: var(--danger) !important;
      font-weight: 700;
    }}
    td.error {{
      color: var(--danger) !important;
      font-weight: 700;
      background: #fff0ee !important;
    }}
    .cell-select {{
      display: block;
      width: 100%;
      min-width: 108px;
      max-width: 108px;
      height: 40px;
      min-height: 40px;
      border: 0;
      outline: none;
      padding: 6px 28px 6px 9px;
      margin: 0 auto;
      color: inherit;
      font: inherit;
      text-align: center;
      text-align-last: center;
      border-radius: 5px;
      background: #fff;
      box-shadow: inset 0 0 0 1px #d9e2ed, inset 0 -2px 0 var(--input-line);
      cursor: pointer;
    }}
    .unit-tool {{
      width: min(100%, 880px);
      min-width: 0;
      padding: 32px;
      background: #fff;
    }}
    .unit-panel {{
      border: 1px solid #dfe8f2;
      border-radius: var(--radius);
      background: #fbfdff;
      overflow: hidden;
    }}
    .unit-panel + .unit-panel {{
      margin-top: 22px;
    }}
    .unit-panel-head {{
      display: flex;
      align-items: center;
      justify-content: space-between;
      min-height: 58px;
      padding: 0 22px;
      border-bottom: 1px solid #e5edf6;
      background: #f6f9fc;
    }}
    .unit-panel-title {{
      font-size: 18px;
      font-weight: 800;
      color: var(--text);
    }}
    .unit-reset {{
      height: 34px;
      padding: 0 16px;
      border: 1px solid #d6e1ee;
      border-radius: 5px;
      background: #fff;
      color: var(--text);
      font-size: 14px;
      font-weight: 700;
      cursor: pointer;
    }}
    .unit-reset:hover {{
      border-color: var(--accent);
      color: var(--accent);
    }}
    .unit-grid {{
      display: grid;
      gap: 0;
      padding: 12px 22px 16px;
    }}
    .unit-row {{
      display: grid;
      grid-template-columns: 130px 76px 240px;
      align-items: center;
      min-height: 60px;
      column-gap: 14px;
      border-bottom: 1px solid #eef3f8;
    }}
    .unit-row:last-child {{
      border-bottom: 0;
    }}
    .unit-label {{
      font-size: 16px;
      font-weight: 800;
      color: var(--text);
      white-space: nowrap;
    }}
    .unit-symbol {{
      font-size: 15px;
      color: var(--muted);
      white-space: nowrap;
    }}
    .unit-field {{
      width: 240px;
      height: 42px;
      border: 0;
      outline: none;
      padding: 8px 12px;
      border-radius: 5px;
      background: #fff;
      color: var(--text);
      font: inherit;
      font-size: 16px;
      font-weight: 700;
      font-variant-numeric: tabular-nums;
      text-align: center;
      box-shadow: inset 0 0 0 1px #d9e2ed, inset 0 -2px 0 var(--input-line);
    }}
    .unit-field:focus {{
      box-shadow: inset 0 0 0 1px var(--accent), inset 0 -2px 0 var(--accent);
    }}
    .status {{
      color: var(--muted);
      font-size: 13px;
      white-space: nowrap;
    }}
    .legend {{
      display: flex;
      align-items: center;
      gap: 12px;
      color: var(--muted);
      font-size: 12px;
      min-height: 34px;
      padding: 0 8px;
      border: 1px solid var(--line);
      border-radius: 6px;
      background: rgba(255, 255, 255, .74);
    }}
    .key {{
      display: inline-flex;
      align-items: center;
      gap: 5px;
    }}
    .swatch {{
      width: 14px;
      height: 14px;
      border-radius: 3px;
      border: 1px solid var(--line);
      display: inline-block;
    }}
    .swatch.input {{
      background: #fff;
      box-shadow: inset 0 -2px 0 var(--input-line);
    }}
    .swatch.formula {{ background: var(--formula); }}
    @media (max-width: 980px) {{
      .app-shell {{
        display: block;
      }}
      .sidebar {{
        position: sticky;
        top: 0;
        z-index: 30;
        height: auto;
        padding: 14px;
      }}
      .brand {{
        padding: 0 2px 12px;
      }}
      .tabs {{
        flex-direction: row;
        overflow-x: auto;
        padding: 12px 0 0;
      }}
      .tab {{
        width: auto;
        flex: 0 0 auto;
        white-space: nowrap;
      }}
      .sidebar-note {{ display: none; }}
      .main-panel {{ padding: 16px 14px 24px; }}
      .topbar {{
        flex-direction: column;
        gap: 12px;
      }}
      .summary {{
        grid-template-columns: repeat(2, minmax(0, 1fr));
      }}
      .actions {{ justify-content: flex-start; }}
      .sheet-shell {{ max-height: calc(100vh - 236px); }}
      .calculator-grid,
      .calculator-grid.fba-layout {{
        min-width: 0;
        grid-template-columns: 1fr;
        grid-template-areas: none;
      }}
      .area-cost,
      .area-ads,
      .area-traffic,
      .area-sales,
      .area-profit,
      .area-settings {{
        grid-area: auto;
      }}
      .unit-tool {{
        min-width: 0;
        max-width: none;
        padding: 18px;
      }}
      .unit-row {{
        grid-template-columns: 98px 54px minmax(150px, 1fr);
      }}
      .unit-field {{
        width: 100%;
      }}
    }}
    @media (max-width: 560px) {{
      .brand h1, .title-block h2 {{ font-size: 18px; }}
      .summary {{ grid-template-columns: 1fr; }}
      .legend {{
        width: 100%;
        justify-content: space-between;
      }}
      .actions button {{
        flex: 1 1 auto;
      }}
    }}
    @media print {{
      body {{ background: #fff; }}
      .sidebar, .topbar, .summary {{ display: none; }}
      .app-shell {{ display: block; }}
      .main-panel {{ padding: 0; }}
      .sheet-shell {{ box-shadow: none; border: 0; max-height: none; }}
    }}
  </style>
</head>
<body>
  <div class="app-shell">
    <aside class="sidebar">
      <div class="brand">
        <h1>AMZ运营计算器2.0</h1>
        <p>利润、广告、尺寸和多店铺投入测算</p>
      </div>
      <nav class="tabs" aria-label="工作表">
        {sheet_tabs}
      </nav>
      <div class="sidebar-note">生成日期：{generated_at}</div>
    </aside>
    <main class="main-panel">
      <header class="topbar">
        <div class="title-block">
          <h2 id="sheetTitle">FBA最低售价及利润</h2>
          <p>带下划线的是可输入项，浅色数据块为公式结果，修改后会自动刷新。</p>
        </div>
        <div class="actions">
          <div class="legend">
            <span class="key"><span class="swatch input"></span>可输入项</span>
            <span class="key"><span class="swatch formula"></span>不可修改</span>
          </div>
          <button type="button" id="resetBtn">恢复默认</button>
          <button type="button" class="primary" id="printBtn">打印 / 另存为 PDF</button>
        </div>
      </header>
      <section class="summary" aria-label="当前工作表信息">
        <div class="metric"><span>当前工作表</span><strong class="small" id="metricSheet">FBA最低售价及利润</strong></div>
        <div class="metric"><span>可输入项</span><strong id="metricInputs">0</strong></div>
        <div class="metric"><span>公式结果</span><strong id="metricFormulas">0</strong></div>
        <div class="metric"><span>错误提示</span><strong id="metricErrors">0</strong></div>
      </section>
      <section class="sheet-shell" id="sheetRoot" aria-live="polite"></section>
    </main>
  </div>
  <script>
    const workbook = {workbook_json};
    let activeSheetIndex = 0;
    let sheetState = new Map();
    let selectState = new Map();
    const displayNames = {{"尺寸": "单位互换"}};

    function displayName(sheet) {{
      return displayNames[sheet.name] || sheet.name;
    }}

    function cloneValues(values) {{
      return Object.fromEntries(Object.entries(values).map(([key, value]) => [key, value]));
    }}

    function currentSheet() {{
      return workbook.sheets[activeSheetIndex];
    }}

    function ensureState(sheet) {{
      if (!sheetState.has(sheet.name)) {{
        sheetState.set(sheet.name, cloneValues(sheet.values));
      }}
      return sheetState.get(sheet.name);
    }}

    function ensureSelectState(sheet) {{
      if (!selectState.has(sheet.name)) {{
        const initial = {{}};
        for (const coord of Object.keys(sheet.selectCells || {{}})) {{
          initial[coord] = "sync";
        }}
        selectState.set(sheet.name, initial);
      }}
      return selectState.get(sheet.name);
    }}

    function isNumericValue(value) {{
      if (value === null || value === undefined || value === "") return false;
      return Number.isFinite(Number(String(value).replace(/,/g, "")));
    }}

    function isNegativeValue(value) {{
      if (!isNumericValue(value)) return false;
      return Number(String(value).replace(/,/g, "")) < 0;
    }}

    function markNumericState(element, value) {{
      element.classList.toggle("numeric", isNumericValue(value));
      element.classList.toggle("negative", isNegativeValue(value));
    }}

    function parseValue(raw) {{
      const trimmed = String(raw ?? "").trim();
      if (trimmed === "") return "";
      const numeric = Number(trimmed.replace(/,/g, ""));
      return Number.isFinite(numeric) ? numeric : trimmed;
    }}

    function formatValue(value) {{
      if (value === null || value === undefined || value === "") return "";
      if (typeof value === "string") return value;
      if (!Number.isFinite(value)) return "#DIV/0!";
      const abs = Math.abs(value);
      const maxDigits = abs >= 1000 ? 2 : 4;
      return value.toLocaleString("zh-CN", {{
        maximumFractionDigits: maxDigits,
        minimumFractionDigits: 0
      }});
    }}

    function cellGetter(state) {{
      return function v(coord) {{
        const value = state[coord.replace(/\\$/g, "")];
        if (value === "" || value === null || value === undefined) return 0;
        const numeric = Number(value);
        return Number.isFinite(numeric) ? numeric : value;
      }};
    }}

    function IF(condition, truthy, falsy) {{
      return condition ? truthy : falsy;
    }}

    function ROUND(value, digits) {{
      const factor = Math.pow(10, Number(digits) || 0);
      return Math.round(Number(value) * factor) / factor;
    }}

    function applySelectCells(sheet) {{
      const state = ensureState(sheet);
      const modes = ensureSelectState(sheet);
      for (const [coord, config] of Object.entries(sheet.selectCells || {{}})) {{
        state[coord] = modes[coord] === "zero" ? 0 : state[config.source];
      }}
    }}

    function cellByCoord(sheet, coord) {{
      return sheet.cells.find((cell) => cell.coord === coord);
    }}

    function renderValueControl(sheet, cell, state) {{
      if (!cell) return document.createTextNode("");
      const selectConfig = sheet.selectCells && sheet.selectCells[cell.coord];
      if (selectConfig) {{
        const modes = ensureSelectState(sheet);
        const select = document.createElement("select");
        select.className = "cell-select";
        select.dataset.coord = cell.coord;
        select.title = `${{selectConfig.syncLabel}}：${{formatValue(state[selectConfig.source])}}`;
        select.setAttribute("aria-label", selectConfig.syncLabel);
        const zeroOption = document.createElement("option");
        zeroOption.value = "zero";
        zeroOption.textContent = selectConfig.zeroLabel;
        const syncOption = document.createElement("option");
        syncOption.value = "sync";
        syncOption.textContent = formatValue(state[selectConfig.source]);
        select.append(zeroOption, syncOption);
        select.value = modes[cell.coord] || "sync";
        markNumericState(select, state[cell.coord]);
        select.addEventListener("change", () => {{
          modes[cell.coord] = select.value;
          recalculate(sheet);
          refreshFormulaCells();
          refreshSelectCells();
        }});
        return select;
      }}
      if (cell.formula) {{
        const valueBox = document.createElement("span");
        valueBox.className = "cell-value";
        valueBox.dataset.coord = cell.coord;
        valueBox.textContent = formatValue(state[cell.coord]);
        valueBox.classList.toggle("emphasis", Boolean(cell.emphasis));
        valueBox.classList.toggle("error", typeof state[cell.coord] === "string" && state[cell.coord].startsWith("#"));
        markNumericState(valueBox, state[cell.coord]);
        return valueBox;
      }}
      if (cell.editable) {{
        const input = document.createElement("input");
        input.className = "cell-input";
        input.dataset.coord = cell.coord;
        input.inputMode = "decimal";
        input.value = state[cell.coord] ?? "";
        markNumericState(input, input.value);
        input.addEventListener("input", () => {{
          state[cell.coord] = parseValue(input.value);
          markNumericState(input, state[cell.coord]);
          recalculate(sheet);
          refreshFormulaCells();
          refreshSelectCells();
        }});
        return input;
      }}
      const valueBox = document.createElement("span");
      valueBox.className = isNumericValue(cell.value) ? "cell-value numeric" : "";
      valueBox.dataset.coord = cell.coord;
      valueBox.textContent = cell.display;
      markNumericState(valueBox, cell.value);
      return valueBox;
    }}

    function addField(card, sheet, labelCoord, valueCoord, unitCoord) {{
      const state = ensureState(sheet);
      const labelCell = cellByCoord(sheet, labelCoord);
      const valueCell = cellByCoord(sheet, valueCoord);
      const unitCell = unitCoord ? cellByCoord(sheet, unitCoord) : null;
      const row = document.createElement("div");
      row.className = unitCoord ? "field-row" : "field-row compact";
      const label = document.createElement("div");
      label.className = "field-label";
      label.textContent = labelCell?.display || "";
      const value = document.createElement("div");
      value.className = "field-control";
      value.appendChild(renderValueControl(sheet, valueCell, state));
      const unit = document.createElement("div");
      unit.className = "field-unit";
      unit.textContent = unitCell?.display || "";
      row.append(label, value);
      if (unitCoord) row.appendChild(unit);
      card.appendChild(row);
    }}

    function addFormulaPair(card, sheet, labelCoord, valueCoord, unitCoord, extraValueCoord = null) {{
      addField(card, sheet, labelCoord, valueCoord, unitCoord);
      if (extraValueCoord) {{
        const row = document.createElement("div");
        row.className = "field-row compact indent-row";
        const spacer = document.createElement("div");
        const value = document.createElement("div");
        value.className = "field-control";
        value.appendChild(renderValueControl(sheet, cellByCoord(sheet, extraValueCoord), ensureState(sheet)));
        row.append(spacer, value);
        card.appendChild(row);
      }}
    }}

    function makeCard(title) {{
      const card = document.createElement("section");
      card.className = "calc-card";
      const heading = document.createElement("h3");
      heading.textContent = title;
      card.appendChild(heading);
      return card;
    }}

    function assignArea(card, area) {{
      card.classList.add(`area-${{area}}`);
      return card;
    }}

    function renderFbaCalculator(sheet, root) {{
      const grid = document.createElement("div");
      grid.className = "calculator-grid fba-layout";
      const base = assignArea(makeCard("成本参数"), "cost");
      [["A1","B1","C1"],["A2","B2","C2"],["A3","B3","C3"],["A4","B4","C4"],["A5","B5","C5"]].forEach((args) => addField(base, sheet, ...args));

      const ads = assignArea(makeCard("广告设置"), "ads");
      [["D1","E1","F1"],["D2","E2","F2"],["D3","E3","F3"],["D4","E4","F4"]].forEach((args) => addField(ads, sheet, ...args));

      const traffic = assignArea(makeCard("流量结果"), "traffic");
      [["G2","H2",null],["G3","H3",null],["G4","H4",null]].forEach((args) => addField(traffic, sheet, ...args));

      const sales = assignArea(makeCard("售价设置"), "sales");
      [["A8","B8","C8"],["A9","B9","C9"],["A10","B10","C10"],["A11","B11","C11"]].forEach((args) => addField(sales, sheet, ...args));
      const divider = document.createElement("hr");
      divider.className = "card-divider";
      const subhead = document.createElement("div");
      subhead.className = "subhead";
      subhead.textContent = "回款与成本";
      sales.append(divider, subhead);
      [["A12","B12","C12"],["A13","B13","C13"],["A14","B14","C14"],["A15","B15","C15"],["A16","B16","C16"]].forEach((args) => addField(sales, sheet, ...args));

      const profit = assignArea(makeCard("利润测算"), "profit");
      [["D7","E7","F7"],["D8","E8","F8"],["D9","E9","F9"],["D10","E10","F10"],["D11","E11","F11"],["D12","E12","F12"]].forEach((args) => addField(profit, sheet, ...args));

      const settings = assignArea(makeCard("汇率与税率"), "settings");
      [["D14","E14",null],["D15","E15",null],["D16","E16",null]].forEach((args) => {{
        if (cellByCoord(sheet, args[0]) || cellByCoord(sheet, args[1])) addField(settings, sheet, ...args);
      }});

      grid.append(base, ads, traffic, sales, profit);
      if (settings.querySelectorAll(".field-row").length) grid.appendChild(settings);
      const noteCell = cellByCoord(sheet, "A18");
      if (noteCell?.display) {{
        const note = document.createElement("div");
        note.className = "calc-note";
        note.style.gridArea = "note";
        note.textContent = noteCell.display;
        grid.appendChild(note);
      }}
      root.replaceChildren(grid);
    }}

    function trimUnitNumber(value) {{
      if (!Number.isFinite(value)) return "";
      const rounded = Math.round((value + Number.EPSILON) * 1000000) / 1000000;
      return String(rounded).replace(/\.?0+$/, "");
    }}

    function renderUnitPanel(title, units, toBase, fromBase) {{
      let syncing = false;
      const panel = document.createElement("section");
      panel.className = "unit-panel";

      const head = document.createElement("div");
      head.className = "unit-panel-head";
      const heading = document.createElement("div");
      heading.className = "unit-panel-title";
      heading.textContent = title;
      const reset = document.createElement("button");
      reset.type = "button";
      reset.className = "unit-reset";
      reset.textContent = "重置";
      head.append(heading, reset);

      const grid = document.createElement("div");
      grid.className = "unit-grid";
      const inputs = {{}};

      units.forEach((unit) => {{
        const row = document.createElement("label");
        row.className = "unit-row";
        const name = document.createElement("span");
        name.className = "unit-label";
        name.textContent = unit.label;
        const symbol = document.createElement("span");
        symbol.className = "unit-symbol";
        symbol.textContent = unit.symbol;
        const input = document.createElement("input");
        input.className = "unit-field";
        input.type = "number";
        input.inputMode = "decimal";
        input.step = "any";
        input.placeholder = "请输入";
        input.dataset.unit = unit.key;
        inputs[unit.key] = input;
        row.append(name, symbol, input);
        grid.appendChild(row);
      }});

      const clearInputs = () => {{
        syncing = true;
        Object.values(inputs).forEach((input) => {{
          input.value = "";
        }});
        syncing = false;
      }};

      const updateGroup = (sourceKey, rawValue) => {{
        if (syncing) return;
        const numericValue = Number(rawValue);
        if (rawValue === "" || !Number.isFinite(numericValue)) {{
          clearInputs();
          return;
        }}
        const baseValue = toBase[sourceKey](numericValue);
        syncing = true;
        Object.entries(inputs).forEach(([key, input]) => {{
          input.value = key === sourceKey ? rawValue : trimUnitNumber(fromBase[key](baseValue));
        }});
        syncing = false;
      }};

      Object.entries(inputs).forEach(([key, input]) => {{
        input.addEventListener("input", () => updateGroup(key, input.value));
      }});
      reset.addEventListener("click", clearInputs);

      panel.append(head, grid);
      return panel;
    }}

    function renderSizeCalculator(sheet, root) {{
      const tool = document.createElement("div");
      tool.className = "unit-tool";
      const lengthUnits = [
        {{ key: "cm", label: "厘米", symbol: "cm" }},
        {{ key: "in", label: "英寸", symbol: "in" }},
      ];
      const weightUnits = [
        {{ key: "kg", label: "公斤", symbol: "kg" }},
        {{ key: "oz", label: "盎司", symbol: "oz" }},
        {{ key: "g", label: "克", symbol: "g" }},
        {{ key: "lb", label: "磅", symbol: "lb" }},
      ];

      tool.append(
        renderUnitPanel(
          "长度单位换算",
          lengthUnits,
          {{
            cm: (value) => value,
            in: (value) => value * 2.54,
          }},
          {{
            cm: (base) => base,
            in: (base) => base / 2.54,
          }},
        ),
        renderUnitPanel(
          "重量单位换算",
          weightUnits,
          {{
            kg: (value) => value,
            oz: (value) => value / 35.27396195,
            g: (value) => value / 1000,
            lb: (value) => value / 2.2046226218,
          }},
          {{
            kg: (base) => base,
            oz: (base) => base * 35.27396195,
            g: (base) => base * 1000,
            lb: (base) => base * 2.2046226218,
          }},
        ),
      );
      root.replaceChildren(tool);
    }}

    function renderAdsBidCalculator(sheet, root) {{
      const grid = document.createElement("div");
      grid.className = "calculator-grid compact-page";
      const base = makeCard("关键词出价");
      addField(base, sheet, "A6", "B6", null);
      const placements = makeCard("广告位模拟");
      [["A2","B2","C2"],["A3","B3","C3"],["A4","B4","C4"]].forEach((args) => addField(placements, sheet, ...args));
      const refs = makeCard("参考结果");
      [["I10","I10",null],["I11","I11",null],["I12","I12",null],["A13","A13",null]].forEach((args) => addField(refs, sheet, ...args));
      grid.append(base, placements, refs);
      root.replaceChildren(grid);
    }}

    function renderAdsBalance(sheet, root) {{
      const grid = document.createElement("div");
      grid.className = "calculator-grid compact-page";
      const traffic = makeCard("流量与转化");
      [["A1","B1",null],["A2","B2",null],["A3","B3",null],["A4","B4",null],["A5","B5",null]].forEach((args) => addField(traffic, sheet, ...args));
      const spend = makeCard("广告花费");
      [["C2","D2",null],["E2","F2",null],["G2","H2",null]].forEach((args) => addField(spend, sheet, ...args));
      const revenue = makeCard("销售与利润");
      [["C3","D3",null],["E3","F3",null],["G3","H3",null],["E4","F4",null],["G4","H4",null]].forEach((args) => addField(revenue, sheet, ...args));
      grid.append(traffic, spend, revenue);
      root.replaceChildren(grid);
    }}

    function renderTemplateLinks(sheet, root) {{
      const grid = document.createElement("div");
      grid.className = "calculator-grid compact-page";
      const card = makeCard("选品模版");
      const linkCell = cellByCoord(sheet, "B5");
      const label = document.createElement("div");
      label.className = "calc-note";
      const a = document.createElement("a");
      a.href = linkCell?.display || "#";
      a.target = "_blank";
      a.rel = "noreferrer";
      a.textContent = linkCell?.display || "";
      label.textContent = `${{cellByCoord(sheet, "A5")?.display || ""}} `;
      label.appendChild(a);
      card.appendChild(label);
      grid.appendChild(card);
      root.replaceChildren(grid);
    }}

    function renderFbmCalculator(sheet, root) {{
      const grid = document.createElement("div");
      grid.className = "calculator-grid";
      const base = makeCard("基础参数");
      [["J15","L15","M15"],["J16","L16","M16"],["J17","L17","M17"],["J18","L18","M18"],["J19","L19","M19"],["J20","L20","M20"]].forEach((args) => addField(base, sheet, ...args));
      const us = makeCard("美国测算");
      [["A1","B1","C1"],["A2","B2","C2"],["A3","B3","C3"],["A4","B4","C4"],["A5","B5","C5"],["A8","B8","C8"],["A9","B9","C9"],["A10","B10","C10"],["A11","B11","C11"],["A12","B12","C12"],["A13","B13","C13"],["A14","B14","C14"],["A16","B16",null]].forEach((args) => addField(us, sheet, ...args));
      const ca = makeCard("加拿大测算");
      [["T1","U1","V1"],["T2","U2","V2"],["T3","U3","V3"],["T4","U4","V4"],["T5","U5","V5"],["T8","U8","V8"],["T9","U9","V9"],["T10","U10","V10"],["T11","U11","V11"],["T12","U12","V12"],["T13","U13","V13"],["T14","U14","V14"],["T16","U16",null]].forEach((args) => addField(ca, sheet, ...args));
      const compare = makeCard("人民币结果");
      [["D2","E2","F2"],["D8","E8","F8"],["D11","E11","F11"],["D12","E12","F12"],["D13","E13","F13"],["D14","E14","F14"],["W2","X2","Y2"],["W8","X8","Y8"],["W11","X11","Y11"],["W12","X12","Y12"],["W13","X13","Y13"],["W14","X14","Y14"]].forEach((args) => addField(compare, sheet, ...args));
      const freight = makeCard("物流价格表");
      [["H3","L3",null],["N3","R3",null],["H6","I6","J6"],["H7","I7","J7"],["H8","I8","J8"],["H9","I9","J9"],["H10","I10","J10"],["N6","O6","P6"],["N7","O7","P7"],["N8","O8","P8"],["N9","O9","P9"],["N10","O10","P10"],["N11","O11","P11"],["N12","O12","P12"],["N13","O13","P13"]].forEach((args) => addField(freight, sheet, ...args));
      grid.append(base, us, ca, compare, freight);
      root.replaceChildren(grid);
    }}

    function renderMultiStoreCalculator(sheet, root) {{
      const grid = document.createElement("div");
      grid.className = "calculator-grid";
      const base = makeCard("基础参数");
      [["A1","B1",null],["A2","B2","C2"],["A3","B3","C3"],["A4","B4","C4"],["D1","E1",null],["D2","E2",null],["D4","E4","F4"],["D5","E5","F5"],["A5","B5","C5"],["D6","E6","F6"],["A6","B6","C6"],["A7","B7","C7"]].forEach((args) => addField(base, sheet, ...args));
      const setup = makeCard("其他支出");
      [["G1","H1","I1"],["G2","H2","I2"],["G3","H3","I3"],["G4","H4","I4"],["G5","H5","I5"],["G6","H6","I6"]].forEach((args) => addField(setup, sheet, ...args));
      const month1 = makeCard("第1个月");
      [["A10","B10",null],["A11","B11","C11"],["A13","B13","C13"]].forEach((args) => addField(month1, sheet, ...args));
      const month2 = makeCard("第2个月");
      [["A16","B16","C16"],["D16","E16","F16"],["A17","B17","C17"],["D17","E17","F17"],["G17","G17","H17"],["A19","B19","C19"],["D19","E19","F19"],["G19","G19","H19"],["A20","B20","C20"],["D20","E20","F20"],["G20","G20","H20"],["A22","B22","C22"],["A24","B24","C24"],["D24","E24","F24"],["D25","E25","F25"]].forEach((args) => addField(month2, sheet, ...args));
      const month3 = makeCard("第3个月及以后");
      [["A27","B27","C27"],["D27","E27","F27"],["D28","E28","F28"],["G28","G28","H28"],["A29","B29","C29"],["D30","E30","F30"],["G30","G30","H30"],["D31","E31","F31"],["G31","G31","H31"],["A33","B33","C33"],["D33","E33","F33"],["A36","B36","C36"],["A37","B37","C37"]].forEach((args) => addField(month3, sheet, ...args));
      grid.append(base, setup, month1, month2, month3);
      const noteCell = cellByCoord(sheet, "A39");
      if (noteCell?.display) {{
        const note = document.createElement("div");
        note.className = "calc-note";
        note.textContent = noteCell.display;
        grid.appendChild(note);
      }}
      root.replaceChildren(grid);
    }}

    function renderCustomSheet(sheet, root) {{
      if (sheet.name === "FBA最低售价及利润" || sheet.name === "欧洲-FBA最低售价及利润") {{
        renderFbaCalculator(sheet, root);
        return true;
      }}
      if (sheet.name === "FBM-北美") {{
        renderFbmCalculator(sheet, root);
        return true;
      }}
      if (sheet.name === "尺寸") {{
        renderSizeCalculator(sheet, root);
        return true;
      }}
      if (sheet.name === "广告竞价模拟") {{
        renderAdsBidCalculator(sheet, root);
        return true;
      }}
      if (sheet.name === "广告平衡") {{
        renderAdsBalance(sheet, root);
        return true;
      }}
      if (sheet.name === "选品模版") {{
        renderTemplateLinks(sheet, root);
        return true;
      }}
      if (sheet.name === "多店铺刷单模式投入计算") {{
        renderMultiStoreCalculator(sheet, root);
        return true;
      }}
      return false;
    }}

    function compileFormula(formula) {{
      if (!formula || formula.includes("#REF!")) {{
        return () => "#REF!";
      }}
      let expression = formula.slice(1)
        .replace(/<>/g, "!=")
        .replace(/\\^/g, "**")
        .replace(/\\bIF\\s*\\(/gi, "IF(")
        .replace(/\\bROUND\\s*\\(/gi, "ROUND(");

      expression = expression.replace(/(?<![A-Z0-9_])\\$?([A-Z]{{1,3}})\\$?(\\d+)\\b/g, (_, col, row) => {{
        return `v("${{col}}${{row}}")`;
      }});
      try {{
        return new Function("v", "IF", "ROUND", `"use strict"; return (${{expression}});`);
      }} catch (error) {{
        return () => "#ERR!";
      }}
    }}

    const compiled = new Map();
    for (const sheet of workbook.sheets) {{
      const formulas = new Map();
      for (const [coord, formula] of Object.entries(sheet.formulas)) {{
        formulas.set(coord, compileFormula(formula));
      }}
      compiled.set(sheet.name, formulas);
    }}

    function recalculate(sheet) {{
      const state = ensureState(sheet);
      const formulas = compiled.get(sheet.name);
      for (let pass = 0; pass < 8; pass += 1) {{
        let changed = false;
        applySelectCells(sheet);
        const v = cellGetter(state);
        for (const [coord, fn] of formulas.entries()) {{
          if (sheet.selectCells && sheet.selectCells[coord]) continue;
          let next;
          try {{
            next = fn(v, IF, ROUND);
          }} catch (error) {{
            next = "#ERR!";
          }}
          if (typeof next === "number" && !Number.isFinite(next)) {{
            next = "#DIV/0!";
          }}
          if (state[coord] !== next) {{
            state[coord] = next;
            changed = true;
          }}
        }}
        if (!changed) break;
      }}
    }}

    function renderSheet(index) {{
      activeSheetIndex = index;
      const sheet = currentSheet();
      recalculate(sheet);
      const state = ensureState(sheet);
      const root = document.getElementById("sheetRoot");
      if (renderCustomSheet(sheet, root)) {{
        document.querySelectorAll(".tab").forEach((tab, tabIndex) => {{
          tab.classList.toggle("is-active", tabIndex === index);
        }});
        updateSheetChrome(sheet);
        return;
      }}
      const table = document.createElement("table");
      table.className = "sheet";

      const colgroup = document.createElement("colgroup");
      for (const col of sheet.colWidths) {{
        const colEl = document.createElement("col");
        colEl.style.width = `${{col.px}}px`;
        if (col.hidden) colEl.style.display = "none";
        colgroup.appendChild(colEl);
      }}
      table.appendChild(colgroup);

      const byRow = new Map();
      for (const cell of sheet.cells) {{
        if (!byRow.has(cell.row)) byRow.set(cell.row, []);
        byRow.get(cell.row).push(cell);
      }}

      const tbody = document.createElement("tbody");
      for (let row = 1; row <= sheet.rows; row += 1) {{
        const tr = document.createElement("tr");
        const rowMeta = sheet.rowHeights[row - 1];
        tr.style.height = `${{rowMeta ? Math.max(42, Math.round(rowMeta.px * 1.15)) : 42}}px`;
        if (rowMeta?.hidden) tr.style.display = "none";
        const cells = byRow.get(row) || [];
        cells.sort((a, b) => a.col - b.col);
        for (const cell of cells) {{
          const td = document.createElement("td");
          td.dataset.coord = cell.coord;
          const colRole = sheet.colWidths[cell.col - 1]?.role;
          if (colRole) td.classList.add(`role-${{colRole}}`);
          if (!cell.value && !cell.formula && !cell.editable && !(sheet.selectCells && sheet.selectCells[cell.coord])) {{
            td.classList.add("blank");
          }}
          if (cell.rowspan > 1) td.rowSpan = cell.rowspan;
          if (cell.colspan > 1) td.colSpan = cell.colspan;
          td.setAttribute("style", cell.style);
          if (String(cell.style).includes("white-space:pre-wrap")) td.classList.add("wrap");
          if (sheet.selectCells && sheet.selectCells[cell.coord]) {{
            td.classList.add("editable", "numeric");
            const config = sheet.selectCells[cell.coord];
            const modes = ensureSelectState(sheet);
            const select = document.createElement("select");
            select.className = "cell-select";
            select.title = `${{config.syncLabel}}：${{formatValue(state[config.source])}}`;
            select.setAttribute("aria-label", config.syncLabel);
            const zeroOption = document.createElement("option");
            zeroOption.value = "zero";
            zeroOption.textContent = config.zeroLabel;
            const syncOption = document.createElement("option");
            syncOption.value = "sync";
            syncOption.textContent = formatValue(state[config.source]);
            select.append(zeroOption, syncOption);
            select.value = modes[cell.coord] || "sync";
            markNumericState(select, state[cell.coord]);
            select.addEventListener("change", () => {{
              modes[cell.coord] = select.value;
              recalculate(sheet);
              refreshFormulaCells();
              refreshSelectCells();
            }});
            td.appendChild(select);
          }} else if (cell.formula) {{
            td.classList.add("formula");
            td.title = cell.formula;
            const formulaValue = state[cell.coord];
            const valueBox = document.createElement("span");
            valueBox.className = "cell-value";
            valueBox.textContent = formatValue(formulaValue);
            valueBox.classList.toggle("emphasis", Boolean(cell.emphasis));
            valueBox.classList.toggle("error", typeof formulaValue === "string" && formulaValue.startsWith("#"));
            markNumericState(valueBox, formulaValue);
            td.appendChild(valueBox);
            markNumericState(td, formulaValue);
          }} else if (cell.editable) {{
            td.classList.add("editable");
            const input = document.createElement("input");
            input.className = "cell-input";
            input.inputMode = "decimal";
            input.value = state[cell.coord] ?? "";
            markNumericState(input, input.value);
            input.addEventListener("input", () => {{
              state[cell.coord] = parseValue(input.value);
              markNumericState(input, state[cell.coord]);
              recalculate(sheet);
              refreshFormulaCells();
              refreshSelectCells();
            }});
            td.appendChild(input);
          }} else {{
            if (typeof cell.value === "string" && /^https?:\\/\\//i.test(cell.value)) {{
              const link = document.createElement("a");
              link.href = cell.value;
              link.target = "_blank";
              link.rel = "noreferrer";
              link.textContent = cell.value;
              td.appendChild(link);
            }} else {{
              td.textContent = cell.display;
              markNumericState(td, cell.value);
            }}
          }}
          tr.appendChild(td);
        }}
        tbody.appendChild(tr);
      }}
      table.appendChild(tbody);
      root.replaceChildren(table);
      document.querySelectorAll(".tab").forEach((tab, tabIndex) => {{
        tab.classList.toggle("is-active", tabIndex === index);
      }});
      updateSheetChrome(sheet);
    }}

    function updateSheetChrome(sheet) {{
      const state = ensureState(sheet);
      const errorCount = Object.keys(sheet.formulas).filter((coord) => {{
        const value = state[coord];
        return typeof value === "string" && value.startsWith("#");
      }}).length;
      document.getElementById("sheetTitle").textContent = displayName(sheet);
      document.getElementById("metricSheet").textContent = displayName(sheet);
      document.getElementById("metricInputs").textContent = String(sheet.cells.filter((cell) => cell.editable).length + Object.keys(sheet.selectCells || {{}}).length);
      document.getElementById("metricFormulas").textContent = String(Object.keys(sheet.formulas).length);
      document.getElementById("metricErrors").textContent = String(errorCount);
    }}

    function refreshFormulaCells() {{
      const sheet = currentSheet();
      const state = ensureState(sheet);
      document.querySelectorAll("td.formula").forEach((td) => {{
        const value = state[td.dataset.coord];
        const valueBox = td.querySelector(".cell-value");
        if (valueBox) {{
          valueBox.textContent = formatValue(value);
          valueBox.classList.toggle("error", typeof value === "string" && value.startsWith("#"));
          markNumericState(valueBox, value);
        }}
        td.classList.toggle("error", typeof value === "string" && value.startsWith("#"));
        markNumericState(td, value);
      }});
      document.querySelectorAll(".cell-value[data-coord]").forEach((valueBox) => {{
        const value = state[valueBox.dataset.coord];
        if (value === undefined) return;
        valueBox.textContent = formatValue(value);
        valueBox.classList.toggle("error", typeof value === "string" && value.startsWith("#"));
        markNumericState(valueBox, value);
      }});
      updateSheetChrome(sheet);
    }}

    function refreshSelectCells() {{
      const sheet = currentSheet();
      const state = ensureState(sheet);
      for (const [coord, config] of Object.entries(sheet.selectCells || {{}})) {{
        const select = document.querySelector(`td[data-coord="${{coord}}"] select`);
        if (!select) continue;
        const syncOption = select.querySelector('option[value="sync"]');
        if (syncOption) syncOption.textContent = formatValue(state[config.source]);
        select.title = `${{config.syncLabel}}：${{formatValue(state[config.source])}}`;
        markNumericState(select, state[coord]);
      }}
    }}

    document.querySelectorAll(".tab").forEach((button, index) => {{
      button.addEventListener("click", () => renderSheet(index));
    }});

    document.getElementById("resetBtn").addEventListener("click", () => {{
      sheetState.delete(currentSheet().name);
      renderSheet(activeSheetIndex);
    }});

    document.getElementById("printBtn").addEventListener("click", () => window.print());

    renderSheet(0);
  </script>
</body>
</html>
"""


def main() -> None:
    if not SOURCE.exists() and not FALLBACK_SOURCE.exists():
        raise SystemExit(f"Missing source file: {SOURCE}")
    data = build_workbook_data()
    html = html_template(data)
    OUTPUT.write_text(html, encoding="utf-8")
    INDEX_OUTPUT.write_text(html, encoding="utf-8")
    print(OUTPUT.resolve())
    print(INDEX_OUTPUT.resolve())


if __name__ == "__main__":
    main()
